"""``DocumentParser`` that reads PDF, Word (.docx) and plain-text files locally with pypdf and python-docx.

Everything happens in memory: no file is written to disk or sent anywhere.
"""

import io
import logging
import zipfile
from pathlib import PurePath

import docx
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from app.application.ports.documents import UnreadableDocumentError, UnsupportedDocumentError
from app.domain.documents import ParsedDocument

logger = logging.getLogger("pebble.parsing")

TEXT_EXTENSIONS = {".txt", ".md"}


class LocalDocumentParser:
    def parse(self, content: bytes, filename: str) -> ParsedDocument:
        path = PurePath(filename)
        extension = path.suffix.lower()
        if extension == ".pdf":
            text, pages = self._pdf(content, filename)
        elif extension == ".docx":
            text, pages = self._docx(content, filename), 1
        elif extension in TEXT_EXTENSIONS:
            text, pages = _decode(content), 1
        else:
            raise UnsupportedDocumentError(filename)
        return ParsedDocument(title=path.name, text=text.strip(), pages=pages)

    @staticmethod
    def _pdf(content: bytes, filename: str) -> tuple[str, int]:
        try:
            reader = PdfReader(io.BytesIO(content))
            if reader.is_encrypted:
                raise UnreadableDocumentError(filename)
            texts = [page.extract_text() or "" for page in reader.pages]
        except UnreadableDocumentError:
            raise
        except (PdfReadError, ValueError, KeyError, TypeError, OSError) as e:
            logger.info("Unreadable PDF %s: %s", filename, type(e).__name__)
            raise UnreadableDocumentError(filename) from e
        return "\n\n".join(t.strip() for t in texts if t.strip()), len(texts)

    @staticmethod
    def _docx(content: bytes, filename: str) -> str:
        try:
            document = docx.Document(io.BytesIO(content))
        except (zipfile.BadZipFile, KeyError, ValueError) as e:
            logger.info("Unreadable DOCX %s: %s", filename, type(e).__name__)
            raise UnreadableDocumentError(filename) from e
        blocks = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                blocks.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(b for b in blocks if b.strip())


def _decode(content: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("latin-1")
